import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Source base presentation (official VIT template already formatted)
BASE_PPTX = r"C:\Users\nisar\.gemini\antigravity\brain\ead474d9-213f-4180-9e2f-f96bf91964e2\.user_uploaded\media_1791228376845.pptx"
DESKTOP = os.path.join(os.path.expanduser('~'), 'OneDrive', 'Desktop')

TEAM1_OUT = os.path.join(DESKTOP, "Review1_OfficialTemplate_Team1_Drone.pptx")
TEAM2_OUT = os.path.join(DESKTOP, "Review1_OfficialTemplate_Team2_Ground_AI.pptx")

def build_team1_presentation():
    # Make a copy of the base PPTX
    shutil.copy(BASE_PPTX, TEAM1_OUT)
    prs = pptx.Presentation(TEAM1_OUT)
    
    # Slide 1: Title & Team Details
    s1 = prs.slides[0]
    for s in s1.shapes:
        if s.has_text_frame:
            t = s.text_frame.text
            if "PROJECT ID:" in t:
                s.text_frame.text = "PROJECT ID:  15377IDP0_  (Assigned by Guide)"
            elif "AI IN STRUCTURAL HEALTH MONITORING" in t:
                s.text_frame.text = "AI IN STRUCTURAL HEALTH MONITORING FOR PREDICTIVE MAINTENANCE\nSUBSYSTEM 1: AUTONOMOUS AERIAL DRONE CARRIER & STANDOFF LIDAR SENSING"
    
    # Update Team 1 table on Slide 1
    tables = [s for s in s1.shapes if s.has_table]
    if tables:
        tbl = tables[0].table
        # Rows: 0: Header, 1: Member 1, 2: Member 2, 3..5: Collaboration
        members = [
            ("1", "Priyam Sharma", "25BAI0159", "SCOPE"),
            ("2", "Ayushman Kaushik", "25BCE0927", "SCOPE"),
            ("3", "[Collaborating Ground Team: Nisarg, Mudit, Aryan]", "Phase 2 Team", "SCOPE"),
            ("4", "-", "-", "-"),
            ("5", "-", "-", "-")
        ]
        for r_idx, row_data in enumerate(members, start=1):
            if r_idx < len(tbl.rows):
                for c_idx, val in enumerate(row_data):
                    if c_idx < len(tbl.columns):
                        tbl.cell(r_idx, c_idx).text = val
    
    # Slide 3: Brief Description of Subsystem 1
    s3 = prs.slides[2]
    for s in s3.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "Domain:" in t:
                s.text_frame.text = "Domain:  Aerial Robotics • ToF LiDAR Standoff Sensing • Edge Point Cloud Ingestion • Wireless Telemetry\nApplication Area:  High-rise building facades, bridge columns, elevated ceilings, and inaccessible structural infrastructure"
            elif "Problem Statement" in t:
                pass
            elif "Structural inspection is largely" in t or "What real-world problem" in t:
                s.text_frame.text = "Manual inspection of elevated building facades and structural pillars requires dangerous scaffolding and ladders. Ground-based manual sensor dragging suffers from severe hand tremors (±3 cm jitter), inconsistent sweep speeds, and blind spots. An aerial carrier is needed for stable, close-range structural profiling."
            elif "Proposed Solution" in t:
                pass
            elif "TF-Luna 1D ToF LiDAR streams" in t or "What are you building" in t:
                s.text_frame.text = "Quadcopter aerial platform carrying a TF-Luna LiDAR on a vibration-dampened 2-axis gimbal. Implements a 1.0 m 'Bat-Inspired' standoff hold to scan surface contours without noise clutter, executing automated raster scan paths and streaming edge-compressed PLY telemetry via ESP32 Wi-Fi UDP to Team 2's ground station."
            elif "Expected Outcome" in t:
                pass
            elif "The current prototype performs" in t or "What will the final system" in t:
                s.text_frame.text = "A functional aerial standoff scanning unit maintaining a fixed 1.0 m distance from walls, streaming high-frequency (100 Hz) range data, performing edge voxel downsampling, and eliminating elevated structural inspection hazards."

    # Slide 4: System Pipeline for Drone Subsystem
    s4 = prs.slides[3]
    for s in s4.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "Proposed Solution" in t:
                s.text_frame.text = "Subsystem 1 Pipeline: Autonomous Aerial Data Acquisition"
            elif "Each stage cross-checked" in t:
                s.text_frame.text = "End-to-end hardware, standoff control, edge compression, and wireless handoff to Team 2"

    # Slide 6: Highlight Box Dragging -> Drone Transition
    s6 = prs.slides[5]
    for s in s6.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "CURRENT ENGINEERING CHALLENGE" in t:
                s.text_frame.text = "CURRENT BENCHMARK: FROM MANUAL BOX DRAGGING TO AERIAL STABILITY"
            elif "During physical scanning" in t:
                s.text_frame.text = "Physical Validation: Initial scans were conducted using our custom enclosure box dragged manually along walls. This confirmed the TF-Luna optical response but revealed manual speed variations and hand-tilt jitter.\n\nDrone Flight Strategy (The Bat Analogy): Mounting the sensor on a dampened gimbal with a 1.0 m standoff hold eliminates manual hand tremor, maintaining precise perpendicular laser alignment throughout vertical and horizontal grid sweeps."

    # Slide 7: Key Techniques in Subsystem 1
    s7 = prs.slides[6]
    techniques_t1 = [
        ("Time-of-Flight LiDAR", "Distance d = ct/2 from pulse round-trip time; TF-Luna (850 nm, 100 Hz, ±1 cm precision)."),
        ("1.0 m Standoff Hold (Bat Strategy)", "Maintains tight 1.0 m standoff to wall to maximize beam density and avoid floor/ceiling background clutter."),
        ("2-Axis Gimbal & Vibration Dampening", "Silicone dampers isolate quadcopter motor high-frequency harmonics, ensuring optical beam stability."),
        ("Raster Grid Flight Path", "Systematic serpentine sweep trajectory ensuring complete overlapping spatial coverage of structural walls."),
        ("Edge PLY Ingestion & Voxel Filtering", "Onboard conversion of raw pulses into 3D .PLY format with voxel grid compression before transmission."),
        ("Wireless UDP Telemetry Streaming", "ESP32 Wi-Fi UDP streaming packet: [timestamp, dist_x, dist_y, dist_z, flux, temp] at 100 Hz (<15 ms latency).")
    ]
    # Update textboxes in Slide 7
    # Note: shapes in slide 7 have technique names and explanations
    tech_idx = 0
    for s in s7.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if t in ["Structural Health Monitoring", "ToF LiDAR Distance Sensing", "Calibration & Signal Filtering", 
                     "Geometric Anomaly Detection", "3D Point-Cloud Visualization", "AI/ML Predictive Maintenance"]:
                if tech_idx < len(techniques_t1):
                    s.text_frame.text = techniques_t1[tech_idx][0]
            elif "What it is:" in t:
                if tech_idx < len(techniques_t1):
                    s.text_frame.text = f"What it is: {techniques_t1[tech_idx][1]}"
                    tech_idx += 1

    # Slide 11: Research Gaps (Aerial Focus)
    s11 = prs.slides[10]
    gaps_t1 = [
        ("G1: Low-Cost UAV Payload Gap", "Commercial aerial LiDAR systems (Velodyne, RIEGL) weigh >500g and cost >₹15 Lakhs. Small quadcopters require ultra-lightweight (<15g) micro-ToF sensors with reliable optical calibration."),
        ("G2: Aerodynamic Vibration & Jitter Gap", "Rotor wash and motor vibrations introduce severe optical noise in micro-LiDAR readings without dedicated mechanical vibration isolation and digital median filtering."),
        ("G3: Close-Range Standoff Hold Gap", "Standard GPS fails indoors and near concrete facades. Autonomous close-range standoff hold (1.0m) requires direct real-time micro-LiDAR distance feedback control loops."),
        ("G4: Edge Telemetry Bandwidth Gap", "Transmitting raw uncompressed 100 Hz multi-point spatial sweeps over Wi-Fi causes packet drops; edge voxel downsampling is required on the carrier."),
        ("G5: Environmental Optical Noise Gap", "Ambient sunlight and reflective surface glares distort low-cost ToF optical flux, requiring dynamic threshold calibration during aerial transit.")
    ]
    g_idx = 0
    for s in s11.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "gap" in t.lower() and len(t) > 20:
                if g_idx < len(gaps_t1):
                    s.text_frame.text = f"{gaps_t1[g_idx][0]}\n{gaps_t1[g_idx][1]}"
                    g_idx += 1

    # Slide 12: Objectives (Aerial Focus)
    s12 = prs.slides[11]
    for s in s12.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "MAIN OBJECTIVE" in t:
                s.text_frame.text = "MAIN OBJECTIVE\nTo develop a lightweight, autonomous aerial drone carrier and standoff LiDAR sensing payload that maintains a stable 1.0 m inspection distance from structural surfaces and streams edge-compressed telemetry for 3D analysis."
            elif "O1" in t:
                pass
            elif "To develop a TF-Luna-based system for non-contact" in t:
                s.text_frame.text = "To engineer a vibration-dampened 2-axis gimbal payload integrating the TF-Luna micro-ToF sensor on a quadcopter airframe."
            elif "To develop calibration and filtering methods to reduce" in t:
                s.text_frame.text = "To implement real-time 1.0 m standoff distance hold control and median filtering to eliminate aerodynamic jitter and motion blur."
            elif "To detect and spatially visualize structural" in t:
                s.text_frame.text = "To develop an onboard edge processing routine that converts distance readings to .PLY format with voxel grid downsampling."
            elif "To develop a monitoring framework that supports" in t:
                s.text_frame.text = "To implement high-frequency (100 Hz) low-latency (<15 ms) wireless UDP telemetry streaming to Team 2's ground station."

    prs.save(TEAM1_OUT)
    print(f"[OK] Successfully built Team 1 presentation: {TEAM1_OUT}")


def build_team2_presentation():
    # Make a copy of the base PPTX
    shutil.copy(BASE_PPTX, TEAM2_OUT)
    prs = pptx.Presentation(TEAM2_OUT)
    
    # Slide 1: Title & Team Details
    s1 = prs.slides[0]
    for s in s1.shapes:
        if s.has_text_frame:
            t = s.text_frame.text
            if "PROJECT ID:" in t:
                s.text_frame.text = "PROJECT ID:  15377IDP0_  (Assigned by Guide)"
            elif "AI IN STRUCTURAL HEALTH MONITORING" in t:
                s.text_frame.text = "AI IN STRUCTURAL HEALTH MONITORING FOR PREDICTIVE MAINTENANCE\nSUBSYSTEM 2: 3D POINT CLOUD RECONSTRUCTION, ARCHITECTURAL METROLOGY & AI DEFECT ANALYTICS"
    
    # Update Team 2 table on Slide 1
    tables = [s for s in s1.shapes if s.has_table]
    if tables:
        tbl = tables[0].table
        members = [
            ("1", "Nisarg Joshi", "25BCE0594", "SCOPE"),
            ("2", "Mudit Gupta", "25BAI0110", "SCOPE"),
            ("3", "Aryan Mithari", "25BAI0115", "SCOPE"),
            ("4", "[Aerial Telemetry Feed: Priyam Sharma & Ayushman Kaushik]", "Phase 1 Team", "SCOPE"),
            ("5", "-", "-", "-")
        ]
        for r_idx, row_data in enumerate(members, start=1):
            if r_idx < len(tbl.rows):
                for c_idx, val in enumerate(row_data):
                    if c_idx < len(tbl.columns):
                        tbl.cell(r_idx, c_idx).text = val
    
    # Slide 3: Brief Description of Subsystem 2
    s3 = prs.slides[2]
    for s in s3.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "Domain:" in t:
                s.text_frame.text = "Domain:  3D Point Cloud Processing • Structural Metrology • AI Anomaly Detection • WebGL Visual Analytics\nApplication Area:  Civil structural inspection, building dimension calculation, defect volumetric quantification, and digital twin monitoring"
            elif "Problem Statement" in t:
                pass
            elif "Structural inspection is largely" in t or "What real-world problem" in t:
                s.text_frame.text = "Raw distance telemetry from sensors is merely 1D time-series numbers without spatial context or actionable engineering insight. Traditional visual inspection cannot quantify structural cavity volumes, subsurface bulges, or crack depths, while expensive 3D CAD modeling software lacks automated defect classification."
            elif "Proposed Solution" in t:
                pass
            elif "TF-Luna 1D ToF LiDAR streams" in t or "What are you building" in t:
                s.text_frame.text = "A Ground Analysis Engine that ingests wireless UDP telemetry from Team 1's aerial carrier, performs spherical-to-Cartesian transformation into millimeter-accurate 3D point clouds (.PLY), automatically calculates architectural metrics (Floor/Wall Area, Volume, Perimeter), localizes structural defects (Cavities, Bulges, Cracks), and classifies geometries via Machine Learning."
            elif "Expected Outcome" in t:
                pass
            elif "The current prototype performs" in t or "What will the final system" in t:
                s.text_frame.text = "An interactive 5-Tab Streamlit Web Dashboard featuring real-time WebGL Three.js point cloud visualization, automated metrology reports (Floor Area 15.12 m², Volume 40.82 m³), automated 3-class defect detection, and continuous predictive maintenance tracking."

    # Slide 4: System Pipeline for Ground Analysis
    s4 = prs.slides[3]
    for s in s4.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "Proposed Solution" in t:
                s.text_frame.text = "Subsystem 2 Pipeline: Telemetry Ingestion, 3D Metrology & AI Defect Engine"
            elif "Each stage cross-checked" in t:
                s.text_frame.text = "From raw UDP stream to Delaunay 2.5D meshing, architectural metrology, and PointNet ML classification"

    # Slide 7: Key Techniques in Subsystem 2
    s7 = prs.slides[6]
    techniques_t2 = [
        ("3D Spatial Reconstruction", "Spherical to Cartesian conversion: X = r*sin(θ)*cos(φ), Y = r*sin(θ)*sin(φ), Z = r*cos(θ) into .PLY format."),
        ("Automated Architectural Metrology", "Automated spatial calculations: Floor Area = W*D, Wall Area = 2(W+D)H, Volume = W*D*H, Perimeter = 2(W+D)."),
        ("Multi-Signature Defect Localization", "Cavities (+Δd deeper penetration), Bulges (-Δd surface expansion), and Cracks (sharp optical flux drop < 100)."),
        ("Delaunay 2.5D & Alpha Shape Meshing", "Reconstructs continuous triangular mesh surfaces from point clouds for FEA stress modeling and defect patch isolation."),
        ("Machine Learning Geometry Classifier", "PointNet / SVM classification to automatically identify planar walls vs. curved columns and categorize defect severity."),
        ("WebGL Three.js Digital Twin", "Interactive browser rendering with directional wall tagging (North/South/East/West) and elevation cross-sections.")
    ]
    tech_idx = 0
    for s in s7.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if t in ["Structural Health Monitoring", "ToF LiDAR Distance Sensing", "Calibration & Signal Filtering", 
                     "Geometric Anomaly Detection", "3D Point-Cloud Visualization", "AI/ML Predictive Maintenance"]:
                if tech_idx < len(techniques_t2):
                    s.text_frame.text = techniques_t2[tech_idx][0]
            elif "What it is:" in t:
                if tech_idx < len(techniques_t2):
                    s.text_frame.text = f"What it is: {techniques_t2[tech_idx][1]}"
                    tech_idx += 1

    # Slide 11: Research Gaps (Ground Analysis Focus)
    s11 = prs.slides[10]
    gaps_t2 = [
        ("G1: Low-Cost Structural Metrology Gap", "Commercial terrestrial laser scanners (TLS) require proprietary software and lack automated architectural area/volume extraction from sparse or low-cost ToF point clouds."),
        ("G2: Geometry vs. Image-Only Inspection Gap", "Computer vision crack detection fails under variable lighting and cannot quantify true structural cavity depths or volumetric concrete spalling."),
        ("G3: Multi-Anomaly Concurrent Detection Gap", "Reviewed methods target a single defect class in isolation (cracks OR cavities); lacks a unified multi-signature engine detecting depth anomalies and optical flux drops simultaneously."),
        ("G4: Manual Thresholding vs. AI Classification Gap", "Static deviation thresholds produce false positives on curved architectural columns; requires machine learning surface classification (PointNet) to adapt to planar vs. curved structures."),
        ("G5: Detection-to-Prediction Gap", "Damage assessment studies stop at one-time defect visualization; lacks time-series repeated-scan alignment for predictive deterioration rate forecasting.")
    ]
    g_idx = 0
    for s in s11.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "gap" in t.lower() and len(t) > 20:
                if g_idx < len(gaps_t2):
                    s.text_frame.text = f"{gaps_t2[g_idx][0]}\n{gaps_t2[g_idx][1]}"
                    g_idx += 1

    # Slide 12: Objectives (Ground Analysis Focus)
    s12 = prs.slides[11]
    for s in s12.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if "MAIN OBJECTIVE" in t:
                s.text_frame.text = "MAIN OBJECTIVE\nTo develop an AI-assisted Ground Station framework that transforms raw aerial distance telemetry into calibrated 3D point clouds, computes automated architectural metrology, and classifies multi-class structural defects."
            elif "O1" in t:
                pass
            elif "To develop a TF-Luna-based system for non-contact" in t:
                s.text_frame.text = "To build a telemetry parser that reconstructs raw UDP packets into normalized 3D .PLY point clouds with directional wall tagging."
            elif "To develop calibration and filtering methods to reduce" in t:
                s.text_frame.text = "To develop automated algorithms for computing building Floor Area (15.12 m²), Wall Area (42.12 m²), Volume (40.82 m³), and Perimeter."
            elif "To detect and spatially visualize structural" in t:
                s.text_frame.text = "To detect, localize, and quantify structural surface anomalies (cavities, bulges, and optical flux drop cracks) in 2D and 3D."
            elif "To develop a monitoring framework that supports" in t:
                s.text_frame.text = "To train a Machine Learning model (PointNet/SVM) for automated geometric surface classification and predictive maintenance tracking."

    prs.save(TEAM2_OUT)
    print(f"[OK] Successfully built Team 2 presentation: {TEAM2_OUT}")

if __name__ == "__main__":
    build_team1_presentation()
    build_team2_presentation()
    print("Both presentations successfully generated directly from the official VIT template!")
